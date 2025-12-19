"""
Strategy para lectura de archivos Excel XLSX.
"""

import logging
from datetime import datetime, date
from decimal import Decimal
from typing import Iterator, List, Optional

from apps.fund.services.accounting.strategies.base_reader import BaseReaderStrategy
from apps.fund.services.accounting.core.data_classes import ParsedRow, ColumnMapping

logger = logging.getLogger(__name__)


class XlsxReaderStrategy(BaseReaderStrategy):
    """Reader para archivos Excel XLSX."""
    
    def __init__(
        self,
        column_mapping: List[ColumnMapping] = None,
        sheet_name: str = None,
        has_header: bool = True,
        skip_rows: int = 0
    ):
        super().__init__(column_mapping)
        self.sheet_name = sheet_name
        self.has_header = has_header
        self.skip_rows = skip_rows
    
    @property
    def supported_extensions(self) -> List[str]:
        return ['.xlsx', '.xls']
    
    def can_handle(self, filename: str) -> bool:
        return any(filename.lower().endswith(ext) for ext in self.supported_extensions)
    
    def read(self, file_content: bytes, encoding: str = 'utf-8') -> Iterator[ParsedRow]:
        """Lee archivo XLSX y genera filas."""
        try:
            import openpyxl
            from io import BytesIO
            
            workbook = openpyxl.load_workbook(
                BytesIO(file_content), 
                read_only=True, 
                data_only=True
            )
            
            sheet = self._get_sheet(workbook)
            
            header_row = None
            row_number = 0
            
            for row_idx, row in enumerate(sheet.iter_rows(values_only=True)):
                if row_idx < self.skip_rows:
                    continue
                
                row_number = row_idx + 1
                row_values = [str(cell) if cell is not None else "" for cell in row]
                
                # Saltar filas vacías
                if all(v.strip() == '' for v in row_values):
                    continue
                
                # Capturar header
                if self.has_header and header_row is None:
                    header_row = [col.strip() for col in row_values]
                    continue
                
                # Generar ParsedRow
                yield self._create_parsed_row(row_values, row_number, header_row, row)
            
            workbook.close()
            
        except ImportError:
            self.add_error(0, 'openpyxl no está instalado')
        except Exception as e:
            logger.exception(f"Error leyendo archivo XLSX: {e}")
            self.add_error(0, f'Error general: {str(e)}')
    
    def _get_sheet(self, workbook):
        """Obtiene la hoja de cálculo."""
        if self.sheet_name and self.sheet_name in workbook.sheetnames:
            return workbook[self.sheet_name]
        
        if self.sheet_name:
            self.add_warning(0, f"Hoja '{self.sheet_name}' no encontrada, usando hoja activa")
        
        return workbook.active
    
    def _create_parsed_row(
        self,
        row_values: List[str],
        row_number: int,
        header_row: Optional[List[str]],
        original_row: tuple
    ) -> ParsedRow:
        """Crea un ParsedRow preservando tipos de Excel."""
        data = {}
        errors = []
        
        for mapping in self.column_mapping:
            value, error = self._extract_value(
                row_values, header_row, original_row, mapping
            )
            data[mapping.model_field] = value
            if error:
                errors.append(error)
        
        return ParsedRow(
            row_number=row_number,
            raw_data=row_values,
            data=data,
            is_valid=len(errors) == 0,
            errors=errors
        )
    
    def _extract_value(
        self,
        row_values: List[str],
        header_row: Optional[List[str]],
        original_row: tuple,
        mapping: ColumnMapping
    ) -> tuple:
        """Extrae valor preservando tipos nativos de Excel."""
        try:
            # Determinar índice
            if header_row and not mapping.file_column.isdigit():
                col_idx = header_row.index(mapping.file_column)
            else:
                col_idx = int(mapping.file_column)
            
            # Valor original de Excel
            original_value = original_row[col_idx] if col_idx < len(original_row) else None
            
            # Preservar tipos de Excel
            if mapping.data_type == "date":
                if isinstance(original_value, (datetime, date)):
                    return original_value if isinstance(original_value, date) else original_value.date(), None
            
            if mapping.data_type == "decimal":
                if isinstance(original_value, (int, float)):
                    return Decimal(str(original_value)), None
            
            # Usar valor string
            value = row_values[col_idx].strip() if col_idx < len(row_values) else ""
            
            if not value and mapping.required:
                return None, f"Campo '{mapping.model_field}' es requerido"
            
            return value if value else mapping.default_value, None
            
        except (ValueError, IndexError):
            if mapping.required:
                return None, f"Columna '{mapping.file_column}' no encontrada"
            return mapping.default_value, None